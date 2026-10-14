from pathlib import Path
import json

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parent
DATA_PATH = ROOT / "certificate_data.json"
BG_PATH = ROOT / "extracted_media" / "image1.png"
OUT_DIR = ROOT / "outputs"
PAGES_DIR = OUT_DIR / "pages"
DOCX_PATH = OUT_DIR / "校级奖状_16张_修改版.docx"
PDF_PATH = OUT_DIR / "校级奖状_16张_修改版.pdf"
CONTACT_PATH = OUT_DIR / "校级奖状_16张_修改版_检查预览.png"

FONT_REG = Path(r"C:\Windows\Fonts\msyh.ttc")
FONT_BOLD = Path(r"C:\Windows\Fonts\msyhbd.ttc")
FONT_HEI = Path(r"C:\Windows\Fonts\simhei.ttf")
FONT_TITLE = Path(r"C:\Windows\Fonts\STXINWEI.TTF")

W, H = 3508, 2480  # A4 landscape at 300 dpi.


def font(path: Path, size: int) -> ImageFont.FreeTypeFont:
    if not path.exists():
        path = FONT_REG
    return ImageFont.truetype(str(path), size)


def text_width(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont) -> int:
    if not text:
        return 0
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[2] - box[0]


def text_height(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont) -> int:
    box = draw.textbbox((0, 0), text, font=fnt)
    return box[3] - box[1]


def wrap_chars(draw: ImageDraw.ImageDraw, text: str, fnt: ImageFont.FreeTypeFont, max_width: int):
    lines = []
    current = ""
    for ch in text:
        trial = current + ch
        if text_width(draw, trial, fnt) <= max_width or not current:
            current = trial
        else:
            lines.append(current)
            current = ch
    if current:
        lines.append(current)
    return lines


def fit_lines(
    draw,
    text,
    font_path,
    start_size,
    min_size,
    box_width,
    box_height,
    max_lines=3,
    line_gap=1.18,
):
    for size in range(start_size, min_size - 1, -2):
        fnt = font(font_path, size)
        lines = wrap_chars(draw, text, fnt, box_width)
        line_h = text_height(draw, "国", fnt)
        total_h = int(line_h * len(lines) * line_gap)
        if len(lines) <= max_lines and total_h <= box_height:
            return fnt, lines, total_h, line_h
    fnt = font(font_path, min_size)
    lines = wrap_chars(draw, text, fnt, box_width)
    line_h = text_height(draw, "国", fnt)
    total_h = int(line_h * len(lines) * line_gap)
    return fnt, lines, total_h, line_h


def draw_centered_box(draw, text, box, fnt, fill, bold=False, max_lines=3, line_gap=1.18):
    x, y, w, h = box
    lines = wrap_chars(draw, text, fnt, w)
    if len(lines) > max_lines:
        lines = lines[: max_lines - 1] + ["".join(lines[max_lines - 1 :])]
    line_h = text_height(draw, "国", fnt)
    total_h = line_h * len(lines) * line_gap
    cur_y = y + (h - total_h) / 2
    for line in lines:
        tw = text_width(draw, line, fnt)
        draw.text((x + (w - tw) / 2, cur_y), line, font=fnt, fill=fill)
        cur_y += line_h * line_gap


def draw_left_box(draw, text, box, fnt, fill, line_gap=1.18):
    x, y, w, h = box
    lines = wrap_chars(draw, text, fnt, w)
    line_h = text_height(draw, "国", fnt)
    total_h = line_h * len(lines) * line_gap
    cur_y = y + max(0, (h - total_h) / 2)
    for line in lines:
        draw.text((x, cur_y), line, font=fnt, fill=fill)
        cur_y += line_h * line_gap


def draw_title(draw):
    title_font = font(FONT_TITLE, 240)
    text = "荣誉证书"
    tw = text_width(draw, text, title_font)
    x = (W - tw) / 2
    y = 386
    draw.text((x + 8, y + 8), text, font=title_font, fill="#D8AA5B")
    draw.text((x, y), text, font=title_font, fill="#8A2D1F")


def make_page(bg, payload, record, index):
    im = bg.copy()
    draw = ImageDraw.Draw(im)

    event_font = font(FONT_REG, 52)
    sentence_font = font(FONT_REG, 68)
    award_font = font(FONT_HEI, 108)
    detail_font = font(FONT_REG, 60)
    detail_bold = font(FONT_BOLD, 60)
    small_font = font(FONT_REG, 54)

    draw_centered_box(draw, payload["event"], (640, 232, 2228, 104), event_font, "#9E6A19", max_lines=1)
    draw_title(draw)

    title_text = record["work_title"]
    title_line_gap = 1.28
    title_font, title_lines, title_total_h, title_line_h = fit_lines(
        draw, title_text, FONT_BOLD, 90, 56, 2680, 410, max_lines=3, line_gap=title_line_gap
    )
    cur_y = 780 + (410 - title_total_h) / 2
    for line in title_lines:
        tw = text_width(draw, line, title_font)
        draw.text(((W - tw) / 2, cur_y), line, font=title_font, fill="#111111")
        cur_y += title_line_h * title_line_gap

    sentence = f'在 {payload["event"]} 中，荣获'
    draw_centered_box(draw, sentence, (500, 1260, 2508, 118), sentence_font, "#222222", max_lines=2)
    draw_centered_box(draw, record["award_display"], (0, 1375, W, 185), award_font, "#8A2D1F", max_lines=1)
    draw_centered_box(draw, "特发此证，以资鼓励。", (0, 1583, W, 104), sentence_font, "#222222", max_lines=1)

    draw_left_box(draw, f'团队成员：{record["members"]}', (660, 1780, 2200, 120), detail_font, "#1A1A1A")
    draw_left_box(draw, f'指导老师：{record["teacher"]}', (660, 1910, 2200, 104), detail_font, "#1A1A1A")

    draw_left_box(draw, f'颁发单位：{payload["issuer"]}', (1960, 2050, 1080, 92), small_font, "#1A1A1A")
    draw_left_box(draw, f'颁发日期：{payload["date"]}', (1960, 2168, 1080, 92), small_font, "#1A1A1A")

    return im


def build_docx(page_paths):
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Mm(297)
    section.page_height = Mm(210)
    section.top_margin = Mm(0)
    section.bottom_margin = Mm(0)
    section.left_margin = Mm(0)
    section.right_margin = Mm(0)
    section.header_distance = Mm(0)
    section.footer_distance = Mm(0)

    normal = doc.styles["Normal"]
    normal.font.size = Pt(1)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1

    for idx, page in enumerate(page_paths):
        if idx:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1
        run = p.add_run()
        run.add_picture(str(page), width=Mm(297), height=Mm(210))

    doc.save(DOCX_PATH)


def build_contact_sheet(page_paths):
    thumbs = []
    thumb_w = 780
    thumb_h = int(thumb_w * H / W)
    label_h = 54
    sheet_w = thumb_w * 4
    sheet_h = (thumb_h + label_h) * 4
    sheet = Image.new("RGB", (sheet_w, sheet_h), "#F6F1EA")
    label_font = font(FONT_BOLD, 34)

    for idx, page in enumerate(page_paths):
        im = Image.open(page).convert("RGB")
        im.thumbnail((thumb_w, thumb_h), Image.Resampling.LANCZOS)
        x = (idx % 4) * thumb_w
        y = (idx // 4) * (thumb_h + label_h)
        sheet.paste(im, (x, y))
        d = ImageDraw.Draw(sheet)
        label = f"第 {idx + 1:02d} 张"
        tw = text_width(d, label, label_font)
        d.text((x + (thumb_w - tw) / 2, y + thumb_h + 8), label, font=label_font, fill="#7D4B1E")
    sheet.save(CONTACT_PATH, quality=95)


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    PAGES_DIR.mkdir(parents=True, exist_ok=True)

    payload = json.loads(DATA_PATH.read_text(encoding="utf-8"))
    records = payload["records"]
    bg = Image.open(BG_PATH).convert("RGB").resize((W, H), Image.Resampling.LANCZOS)

    page_paths = []
    pdf_images = []
    for idx, record in enumerate(records, start=1):
        page = make_page(bg, payload, record, idx)
        page_path = PAGES_DIR / f"certificate_{idx:02d}.png"
        page.save(page_path, quality=95)
        page_paths.append(page_path)
        pdf_images.append(page)

    build_docx(page_paths)
    pdf_images[0].save(PDF_PATH, "PDF", resolution=300, save_all=True, append_images=pdf_images[1:])
    build_contact_sheet(page_paths)

    print(f"Created {len(page_paths)} certificate pages")
    print(DOCX_PATH)
    print(PDF_PATH)
    print(CONTACT_PATH)


if __name__ == "__main__":
    main()
